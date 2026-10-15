import datetime as dt
import hmac
import io
import json
import os
import re
import sqlite3
import time
from functools import wraps

from flask import Flask, g, jsonify, request, send_file, send_from_directory

import catalog as C

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("DB_PATH", "/data/nearmiss.db")
ADMIN_PIN = os.environ.get("ADMIN_PIN", "change-me")
DEFAULT_PHARMACIES = [p.strip() for p in os.environ.get("PHARMACIES", "DS Gordon Town").split(",") if p.strip()]

app = Flask(__name__, static_folder=os.path.join(BASE, "static"), static_url_path="")
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0

TYPE_CODES = {t[0] for t in C.ERROR_TYPES}
TIME_CODES = {t[0] for t in C.TIMES}
PERSONNEL_KEYS = {k for k, _ in C.PERSONNEL}
REPORTED_KEYS = {k for k, _ in C.REPORTED_BY}
CAUSE_KEYS = {k for k, _ in C.CAUSES}
ACTION_KEYS = {k for k, _ in C.ACTIONS}
REC_KEYS = {k for k, _ in C.RECS}
MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")

SCHEMA = """
CREATE TABLE IF NOT EXISTS pharmacies(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  name TEXT NOT NULL UNIQUE COLLATE NOCASE,
  active INTEGER NOT NULL DEFAULT 1);
CREATE TABLE IF NOT EXISTS drugs(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  name TEXT NOT NULL,
  name_key TEXT NOT NULL UNIQUE,
  active INTEGER NOT NULL DEFAULT 1,
  created_at TEXT NOT NULL DEFAULT (datetime('now')));
CREATE TABLE IF NOT EXISTS records(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  pharmacy_id INTEGER NOT NULL REFERENCES pharmacies(id),
  date TEXT NOT NULL,
  time_of_day TEXT NOT NULL,
  staff_pharmacists INTEGER NOT NULL DEFAULT 0,
  staff_technicians INTEGER NOT NULL DEFAULT 0,
  staff_assistants INTEGER NOT NULL DEFAULT 0,
  staff_interns INTEGER NOT NULL DEFAULT 0,
  personnel TEXT NOT NULL DEFAULT '[]',
  personnel_names TEXT NOT NULL DEFAULT '',
  reported_by_role TEXT NOT NULL,
  reported_by_name TEXT NOT NULL DEFAULT '',
  drug_id INTEGER NOT NULL REFERENCES drugs(id),
  error_type TEXT NOT NULL,
  causes TEXT NOT NULL DEFAULT '[]',
  considerations TEXT NOT NULL DEFAULT '[]',
  actions TEXT NOT NULL DEFAULT '[]',
  action_other TEXT NOT NULL DEFAULT '',
  adverse_discussed INTEGER NOT NULL DEFAULT 0,
  notes TEXT NOT NULL DEFAULT '',
  created_at TEXT NOT NULL DEFAULT (datetime('now')),
  updated_at TEXT NOT NULL DEFAULT (datetime('now')));
CREATE INDEX IF NOT EXISTS idx_records_date ON records(date);
CREATE INDEX IF NOT EXISTS idx_records_pharm ON records(pharmacy_id, date);
CREATE TABLE IF NOT EXISTS monthly(
  pharmacy_id INTEGER NOT NULL REFERENCES pharmacies(id),
  month TEXT NOT NULL,
  submitted_by TEXT NOT NULL DEFAULT '',
  department TEXT NOT NULL DEFAULT '',
  PRIMARY KEY(pharmacy_id, month));
CREATE TABLE IF NOT EXISTS monthly_recs(
  pharmacy_id INTEGER NOT NULL REFERENCES pharmacies(id),
  month TEXT NOT NULL,
  key TEXT NOT NULL,
  count INTEGER NOT NULL DEFAULT 0,
  PRIMARY KEY(pharmacy_id, month, key));
"""


def norm_name(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def init_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.execute("PRAGMA journal_mode=WAL")
    con.executescript(SCHEMA)
    if con.execute("SELECT COUNT(*) FROM drugs").fetchone()[0] == 0:
        with open(os.path.join(BASE, "drugs_seed.json"), encoding="utf-8") as f:
            names = json.load(f)
        con.executemany(
            "INSERT OR IGNORE INTO drugs(name, name_key) VALUES (?, ?)",
            [(norm_name(n), norm_name(n).upper()) for n in names if norm_name(n)],
        )
    if con.execute("SELECT COUNT(*) FROM pharmacies").fetchone()[0] == 0:
        con.executemany("INSERT OR IGNORE INTO pharmacies(name) VALUES (?)", [(p,) for p in DEFAULT_PHARMACIES])
    con.commit()
    con.close()


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH, timeout=30)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys=ON")
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    con = g.pop("db", None)
    if con is not None:
        con.close()


def admin_required(fn):
    @wraps(fn)
    def wrapper(*a, **kw):
        pin = request.headers.get("X-Admin-Pin", "")
        if not hmac.compare_digest(pin.encode(), ADMIN_PIN.encode()):
            time.sleep(0.5)  # slows down PIN guessing
            return jsonify(error="That PIN isn't right."), 403
        return fn(*a, **kw)

    return wrapper


def err(msg, code=400):
    return jsonify(error=msg), code


# ---------------------------------------------------------------- pages

@app.get("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.get("/healthz")
def healthz():
    db().execute("SELECT 1")
    return "ok"


# ---------------------------------------------------------------- config

@app.get("/api/config")
def config():
    pharmacies = [dict(r) for r in db().execute("SELECT id, name FROM pharmacies WHERE active=1 ORDER BY name")]
    return jsonify(
        pharmacies=pharmacies,
        errorTypes=[{"code": c, "name": n, "report": r} for c, n, r in C.ERROR_TYPES],
        times=[{"code": c, "name": n, "range": r} for c, n, r in C.TIMES],
        staffing=[{"key": k, "label": l} for k, l in C.STAFFING],
        personnel=[{"key": k, "label": l} for k, l in C.PERSONNEL],
        reportedBy=[{"key": k, "label": l} for k, l in C.REPORTED_BY],
        causes=[{"key": k, "label": l} for k, l in C.CAUSES],
        suggestedCauses=C.SUGGESTED_CAUSES,
        consider=C.CONSIDER,
        actions=[{"key": k, "label": l} for k, l in C.ACTIONS],
        recs=[{"key": k, "label": l} for k, l in C.RECS],
        defaultDepartment=C.DEFAULT_DEPARTMENT,
    )


@app.get("/api/drugs")
def drugs():
    rows = db().execute("SELECT id, name FROM drugs WHERE active=1 ORDER BY name COLLATE NOCASE").fetchall()
    return jsonify([[r["id"], r["name"]] for r in rows])


# ---------------------------------------------------------------- admin

@app.post("/api/admin/verify")
@admin_required
def admin_verify():
    return jsonify(ok=True)


@app.get("/api/admin/drugs")
@admin_required
def admin_drugs():
    rows = db().execute(
        """SELECT d.id, d.name, d.active, (SELECT COUNT(*) FROM records r WHERE r.drug_id=d.id) AS uses
           FROM drugs d ORDER BY d.name COLLATE NOCASE"""
    ).fetchall()
    return jsonify([[r["id"], r["name"], r["active"], r["uses"]] for r in rows])


@app.post("/api/admin/drugs")
@admin_required
def admin_add_drugs():
    data = request.get_json(silent=True) or {}
    names = data.get("names") or []
    if isinstance(names, str):
        names = names.splitlines()
    names = [norm_name(n) for n in names if norm_name(n)]
    if not names:
        return err("Enter at least one drug name.")
    if any(len(n) > 200 for n in names):
        return err("Drug names must be 200 characters or fewer.")
    con = db()
    added, restored, skipped = 0, 0, []
    for n in names:
        key = n.upper()
        row = con.execute("SELECT id, active FROM drugs WHERE name_key=?", (key,)).fetchone()
        if row is None:
            con.execute("INSERT INTO drugs(name, name_key) VALUES (?, ?)", (n, key))
            added += 1
        elif not row["active"]:
            con.execute("UPDATE drugs SET active=1 WHERE id=?", (row["id"],))
            restored += 1
        else:
            skipped.append(n)
    con.commit()
    return jsonify(added=added, restored=restored, skipped=skipped)


@app.patch("/api/admin/drugs/<int:did>")
@admin_required
def admin_patch_drug(did):
    data = request.get_json(silent=True) or {}
    con = db()
    row = con.execute("SELECT * FROM drugs WHERE id=?", (did,)).fetchone()
    if not row:
        return err("Drug not found.", 404)
    name, active = row["name"], row["active"]
    if "name" in data:
        name = norm_name(data["name"])
        if not name or len(name) > 200:
            return err("Enter a drug name (200 characters or fewer).")
        clash = con.execute("SELECT id FROM drugs WHERE name_key=? AND id<>?", (name.upper(), did)).fetchone()
        if clash:
            return err("Another drug already has that name.")
    if "active" in data:
        active = 1 if data["active"] else 0
    con.execute("UPDATE drugs SET name=?, name_key=?, active=? WHERE id=?", (name, name.upper(), active, did))
    con.commit()
    return jsonify(ok=True)


@app.get("/api/admin/pharmacies")
@admin_required
def admin_pharmacies():
    rows = db().execute(
        """SELECT p.id, p.name, p.active, (SELECT COUNT(*) FROM records r WHERE r.pharmacy_id=p.id) AS uses
           FROM pharmacies p ORDER BY p.name"""
    ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.post("/api/admin/pharmacies")
@admin_required
def admin_add_pharmacy():
    name = norm_name((request.get_json(silent=True) or {}).get("name"))
    if not name or len(name) > 120:
        return err("Enter a pharmacy name.")
    con = db()
    row = con.execute("SELECT id, active FROM pharmacies WHERE name=?", (name,)).fetchone()
    if row:
        if row["active"]:
            return err("That pharmacy is already on the list.")
        con.execute("UPDATE pharmacies SET active=1 WHERE id=?", (row["id"],))
    else:
        con.execute("INSERT INTO pharmacies(name) VALUES (?)", (name,))
    con.commit()
    return jsonify(ok=True)


@app.patch("/api/admin/pharmacies/<int:pid>")
@admin_required
def admin_patch_pharmacy(pid):
    data = request.get_json(silent=True) or {}
    con = db()
    row = con.execute("SELECT * FROM pharmacies WHERE id=?", (pid,)).fetchone()
    if not row:
        return err("Pharmacy not found.", 404)
    name, active = row["name"], row["active"]
    if "name" in data:
        name = norm_name(data["name"])
        if not name or len(name) > 120:
            return err("Enter a pharmacy name.")
        clash = con.execute("SELECT id FROM pharmacies WHERE name=? AND id<>?", (name, pid)).fetchone()
        if clash:
            return err("Another pharmacy already has that name.")
    if "active" in data:
        active = 1 if data["active"] else 0
    con.execute("UPDATE pharmacies SET name=?, active=? WHERE id=?", (name, active, pid))
    con.commit()
    return jsonify(ok=True)


# ---------------------------------------------------------------- records

def clean_record(data):
    """Validate a posted record; returns (values dict, error message)."""
    con = db()
    v = {}
    try:
        v["pharmacy_id"] = int(data.get("pharmacy_id"))
    except (TypeError, ValueError):
        return None, "Choose the pharmacy."
    if not con.execute("SELECT 1 FROM pharmacies WHERE id=?", (v["pharmacy_id"],)).fetchone():
        return None, "Choose the pharmacy."

    d = str(data.get("date") or "")
    try:
        dt.date.fromisoformat(d)
    except ValueError:
        return None, "Enter the date of the near miss."
    v["date"] = d

    if data.get("time_of_day") not in TIME_CODES:
        return None, "Choose the time of day."
    v["time_of_day"] = data["time_of_day"]

    for key, _label in C.STAFFING:
        try:
            n = int(data.get(key) or 0)
        except (TypeError, ValueError):
            return None, "Staffing numbers must be whole numbers."
        if not 0 <= n <= 99:
            return None, "Staffing numbers must be between 0 and 99."
        v[key] = n

    personnel = [p for p in (data.get("personnel") or []) if p in PERSONNEL_KEYS]
    if not personnel:
        return None, "Select the personnel involved."
    v["personnel"] = json.dumps(sorted(set(personnel)))
    v["personnel_names"] = norm_name(data.get("personnel_names"))[:200]

    if data.get("reported_by_role") not in REPORTED_KEYS:
        return None, "Choose who reported the near miss."
    v["reported_by_role"] = data["reported_by_role"]
    v["reported_by_name"] = norm_name(data.get("reported_by_name"))[:120]

    try:
        v["drug_id"] = int(data.get("drug_id"))
    except (TypeError, ValueError):
        return None, "Choose the drug involved."
    if not con.execute("SELECT 1 FROM drugs WHERE id=?", (v["drug_id"],)).fetchone():
        return None, "Choose the drug involved."

    if data.get("error_type") not in TYPE_CODES:
        return None, "Choose the type of error."
    v["error_type"] = data["error_type"]

    causes = [c for c in (data.get("causes") or []) if c in CAUSE_KEYS]
    if not causes:
        return None, "Select at least one possible cause."
    v["causes"] = json.dumps(sorted(set(causes), key=[k for k, _ in C.CAUSES].index))

    allowed = set(C.CONSIDER.get(v["error_type"], []))
    v["considerations"] = json.dumps([c for c in (data.get("considerations") or []) if c in allowed])

    actions = [a for a in (data.get("actions") or []) if a in ACTION_KEYS]
    if not actions:
        return None, "Select the action taken."
    v["actions"] = json.dumps(sorted(set(actions), key=[k for k, _ in C.ACTIONS].index))
    v["action_other"] = norm_name(data.get("action_other"))[:300] if "other_action" in actions else ""

    v["adverse_discussed"] = 1 if data.get("adverse_discussed") else 0
    v["notes"] = str(data.get("notes") or "").strip()[:2000]
    return v, None


def row_to_record(r):
    d = dict(r)
    for k in ("personnel", "causes", "considerations", "actions"):
        d[k] = json.loads(d[k] or "[]")
    d["adverse_discussed"] = bool(d["adverse_discussed"])
    return d


RECORD_SELECT = """SELECT r.*, d.name AS drug_name, p.name AS pharmacy_name
                   FROM records r JOIN drugs d ON d.id=r.drug_id JOIN pharmacies p ON p.id=r.pharmacy_id"""


@app.get("/api/records")
def list_records():
    where, args = [], []
    if request.args.get("pharmacy"):
        where.append("r.pharmacy_id=?")
        args.append(request.args["pharmacy"])
    month = request.args.get("month", "")
    if MONTH_RE.match(month):
        where.append("substr(r.date,1,7)=?")
        args.append(month)
    if request.args.get("type") in TYPE_CODES:
        where.append("r.error_type=?")
        args.append(request.args["type"])
    for tok in request.args.get("q", "").split()[:6]:
        like = f"%{tok}%"
        where.append("(d.name LIKE ? OR r.notes LIKE ? OR r.personnel_names LIKE ? OR r.reported_by_name LIKE ?)")
        args += [like, like, like, like]
    clause = (" WHERE " + " AND ".join(where)) if where else ""
    con = db()
    total = con.execute(
        "SELECT COUNT(*) FROM records r JOIN drugs d ON d.id=r.drug_id" + clause, args
    ).fetchone()[0]
    try:
        limit = max(1, min(200, int(request.args.get("limit", 50))))
        offset = max(0, int(request.args.get("offset", 0)))
    except ValueError:
        limit, offset = 50, 0
    rows = con.execute(
        RECORD_SELECT + clause + " ORDER BY r.date DESC, r.id DESC LIMIT ? OFFSET ?", args + [limit, offset]
    ).fetchall()
    return jsonify(total=total, items=[row_to_record(r) for r in rows])


@app.get("/api/records/<int:rid>")
def get_record(rid):
    row = db().execute(RECORD_SELECT + " WHERE r.id=?", (rid,)).fetchone()
    if not row:
        return err("That record no longer exists.", 404)
    return jsonify(row_to_record(row))


@app.post("/api/records")
def create_record():
    v, e = clean_record(request.get_json(silent=True) or {})
    if e:
        return err(e)
    cols = ",".join(v)
    cur = db().execute(f"INSERT INTO records({cols}) VALUES ({','.join('?' * len(v))})", list(v.values()))
    db().commit()
    return jsonify(id=cur.lastrowid), 201


@app.put("/api/records/<int:rid>")
def update_record(rid):
    con = db()
    if not con.execute("SELECT 1 FROM records WHERE id=?", (rid,)).fetchone():
        return err("Record not found.", 404)
    v, e = clean_record(request.get_json(silent=True) or {})
    if e:
        return err(e)
    sets = ",".join(f"{k}=?" for k in v)
    con.execute(f"UPDATE records SET {sets}, updated_at=datetime('now') WHERE id=?", list(v.values()) + [rid])
    con.commit()
    return jsonify(id=rid)


@app.delete("/api/records/<int:rid>")
@admin_required
def delete_record(rid):
    con = db()
    con.execute("DELETE FROM records WHERE id=?", (rid,))
    con.commit()
    return jsonify(ok=True)


@app.get("/api/months")
def months():
    rows = db().execute("SELECT DISTINCT substr(date,1,7) AS m FROM records ORDER BY m DESC").fetchall()
    return jsonify([r["m"] for r in rows])


# ---------------------------------------------------------------- report

def compute_report(month, pharmacy_id):
    con = db()
    where, args = ["substr(r.date,1,7)=?"], [month]
    pharmacy = None
    if pharmacy_id:
        pharmacy = con.execute("SELECT id, name FROM pharmacies WHERE id=?", (pharmacy_id,)).fetchone()
        if not pharmacy:
            return None
        where.append("r.pharmacy_id=?")
        args.append(pharmacy_id)
    rows = [row_to_record(r) for r in con.execute(
        RECORD_SELECT + " WHERE " + " AND ".join(where) + " ORDER BY r.date, r.id", args)]

    def tally(keys, getter):
        counts = {k: 0 for k in keys}
        for r in rows:
            for k in getter(r):
                if k in counts:
                    counts[k] += 1
        return counts

    t_types = tally([c for c, *_ in C.ERROR_TYPES], lambda r: [r["error_type"]])
    t_times = tally([c for c, *_ in C.TIMES], lambda r: [r["time_of_day"]])
    t_pers = tally([k for k, _ in C.PERSONNEL], lambda r: r["personnel"])
    t_rep = tally([k for k, _ in C.REPORTED_BY], lambda r: [r["reported_by_role"]])
    t_cause = tally([k for k, _ in C.CAUSES], lambda r: r["causes"])
    t_act = tally([k for k, _ in C.ACTIONS], lambda r: r["actions"])

    drug_counts = {}
    for r in rows:
        e = drug_counts.setdefault(r["drug_id"], {"id": r["drug_id"], "name": r["drug_name"], "count": 0})
        e["count"] += 1
    top_drugs = sorted(drug_counts.values(), key=lambda e: (-e["count"], e["name"].lower()))[:20]

    # Monthly figures typed in on the report page
    if pharmacy_id:
        recs = {r["key"]: r["count"] for r in con.execute(
            "SELECT key, count FROM monthly_recs WHERE pharmacy_id=? AND month=?", (pharmacy_id, month))}
        m = con.execute("SELECT submitted_by, department FROM monthly WHERE pharmacy_id=? AND month=?",
                        (pharmacy_id, month)).fetchone()
        meta = {"submitted_by": m["submitted_by"] if m else "",
                "department": m["department"] if m else C.DEFAULT_DEPARTMENT}
    else:
        recs = {r["key"]: r["count"] for r in con.execute(
            "SELECT key, SUM(count) AS count FROM monthly_recs WHERE month=? GROUP BY key", (month,))}
        meta = None

    by_pharmacy = None
    if not pharmacy_id:
        by_pharmacy = [dict(r) for r in con.execute(
            """SELECT p.id, p.name, COUNT(r.id) AS count FROM pharmacies p
               LEFT JOIN records r ON r.pharmacy_id=p.id AND substr(r.date,1,7)=?
               WHERE p.active=1 OR r.id IS NOT NULL GROUP BY p.id ORDER BY p.name""", (month,))]

    return {
        "month": month,
        "pharmacy": dict(pharmacy) if pharmacy else None,
        "total": len(rows),
        "types": [{"code": c, "label": rep, "count": t_types[c]} for c, _n, rep in C.ERROR_TYPES],
        "times": [{"code": c, "label": f"{n} ({rg})", "count": t_times[c]} for c, n, rg in C.TIMES],
        "personnel": [{"key": k, "label": l, "count": t_pers[k]} for k, l in C.PERSONNEL],
        "reportedBy": [{"key": k, "label": l, "count": t_rep[k]} for k, l in C.REPORTED_BY],
        "causes": [{"key": k, "label": l, "count": t_cause[k]} for k, l in C.CAUSES],
        "actions": [{"key": k, "label": l, "count": t_act[k]} for k, l in C.ACTIONS],
        "adverseDiscussed": sum(1 for r in rows if r["adverse_discussed"]),
        "drugs": top_drugs,
        "recs": [{"key": k, "label": l, "count": int(recs.get(k, 0))} for k, l in C.RECS],
        "meta": meta,
        "byPharmacy": by_pharmacy,
        "_rows": rows,
    }


def report_args():
    month = request.args.get("month", "")
    if not MONTH_RE.match(month):
        return None, None, "Choose a month."
    pid = request.args.get("pharmacy") or None
    if pid is not None:
        try:
            pid = int(pid)
        except ValueError:
            return None, None, "Unknown pharmacy."
    return month, pid, None


@app.get("/api/report")
def report():
    month, pid, e = report_args()
    if e:
        return err(e)
    rep = compute_report(month, pid)
    if rep is None:
        return err("Unknown pharmacy.", 404)
    rep.pop("_rows")
    return jsonify(rep)


@app.put("/api/report/meta")
def report_meta():
    data = request.get_json(silent=True) or {}
    month = data.get("month", "")
    try:
        pid = int(data.get("pharmacy_id"))
    except (TypeError, ValueError):
        return err("Choose a single pharmacy to save these details.")
    if not MONTH_RE.match(month):
        return err("Choose a month.")
    con = db()
    if not con.execute("SELECT 1 FROM pharmacies WHERE id=?", (pid,)).fetchone():
        return err("Unknown pharmacy.", 404)
    if "submitted_by" in data or "department" in data:
        cur = con.execute("SELECT submitted_by, department FROM monthly WHERE pharmacy_id=? AND month=?", (pid, month)).fetchone()
        sb = norm_name(data.get("submitted_by", cur["submitted_by"] if cur else ""))[:120]
        dep = norm_name(data.get("department", cur["department"] if cur else C.DEFAULT_DEPARTMENT))[:160]
        con.execute(
            """INSERT INTO monthly(pharmacy_id, month, submitted_by, department) VALUES (?,?,?,?)
               ON CONFLICT(pharmacy_id, month) DO UPDATE SET submitted_by=excluded.submitted_by, department=excluded.department""",
            (pid, month, sb, dep))
    for key, val in (data.get("recs") or {}).items():
        if key not in REC_KEYS:
            continue
        try:
            n = max(0, min(9999, int(val or 0)))
        except (TypeError, ValueError):
            continue
        con.execute(
            """INSERT INTO monthly_recs(pharmacy_id, month, key, count) VALUES (?,?,?,?)
               ON CONFLICT(pharmacy_id, month, key) DO UPDATE SET count=excluded.count""",
            (pid, month, key, n))
    con.commit()
    return jsonify(ok=True)


# ---------------------------------------------------------------- Excel export

def build_workbook(rep):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary Report"
    thin = Side(style="thin", color="8A9A9A")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="CFE3E3")
    soft = PatternFill("solid", fgColor="EEF4F4")

    def put(ref, val, bold=False, fill=None, border=True, align=None, size=10, wrap=False):
        c = ws[ref]
        c.value = val
        c.font = Font(name="Arial", size=size, bold=bold)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if fill:
            c.fill = fill
        if border:
            c.border = box
        return c

    def merge_box(rng):
        ws.merge_cells(rng)
        for row in ws[rng]:
            for c in row:
                c.border = box

    month_label = dt.date(int(rep["month"][:4]), int(rep["month"][5:]), 1).strftime("%B %Y")
    pharm = rep["pharmacy"]["name"] if rep["pharmacy"] else "All pharmacies"
    meta = rep["meta"] or {}

    ws.merge_cells("A1:H1")
    put("A1", "PHARMACY SERVICES NEAR MISS SUMMARY REPORT", bold=True, size=14, border=False, align="center")
    ws.row_dimensions[1].height = 26

    info = [("NAME OF PHARMACY:", pharm, "NAME OF DEPARTMENT:", meta.get("department", "")),
            ("SUBMITTED BY:", meta.get("submitted_by", ""), None, None),
            ("MONTH:", month_label, None, None),
            ("DATE:", dt.date.today().strftime("%d %B %Y"), None, None)]
    for i, (a, b, e, f) in enumerate(info):
        r = 3 + i
        put(f"A{r}", a, bold=True, fill=soft)
        ws.merge_cells(f"B{r}:C{r}")
        put(f"B{r}", b)
        ws[f"C{r}"].border = box
        if e:
            put(f"E{r}", e, bold=True, fill=soft)
            ws.merge_cells(f"F{r}:H{r}")
            put(f"F{r}", f)
            for col in "GH":
                ws[f"{col}{r}"].border = box

    put("A8", "Table 1: Pharmacy Near Miss Error Report", bold=True, border=False)
    put("E8", "Table 2", bold=True, border=False)

    # ---- left block (A:B)
    r = 9

    def section(title, total):
        nonlocal r
        put(f"A{r}", title, bold=True, fill=head)
        put(f"B{r}", total, bold=True, fill=head, align="center")
        r += 1

    def line(label, n):
        nonlocal r
        put(f"A{r}", label, wrap=True)
        put(f"B{r}", n, align="center")
        r += 1

    put("A9", "", fill=head)
    put("B9", "Total", bold=True, fill=head, align="center")
    r = 10
    section("Number of errors for the month", rep["total"])
    section("Error code / type of error", rep["total"])
    for t in rep["types"]:
        line(t["label"], t["count"])
    section("Time of day", rep["total"])
    for t in rep["times"]:
        line(t["label"], t["count"])
    section("Personnel involved", sum(p["count"] for p in rep["personnel"]))
    for p in rep["personnel"]:
        line(p["label"], p["count"])
    section("Reported by", rep["total"])
    for p in rep["reportedBy"]:
        line(p["label"], p["count"])
    section("Potential adverse event discussed", rep["adverseDiscussed"])
    left_end = r

    # ---- right block, causes (E:F)
    put("E9", "Possible causes (non-exhaustive)", bold=True, fill=head)
    put("F9", sum(c["count"] for c in rep["causes"]), bold=True, fill=head, align="center")
    rr = 10
    for c in rep["causes"]:
        put(f"E{rr}", c["label"], wrap=True)
        put(f"F{rr}", c["count"], align="center")
        rr += 1
    put(f"E{rr}", "Total", bold=True, fill=soft)
    put(f"F{rr}", sum(c["count"] for c in rep["causes"]), bold=True, fill=soft, align="center")
    causes_end = rr + 1

    # ---- right block, actions and recommendations (G:H)
    put("G9", "Action taken", bold=True, fill=head)
    put("H9", sum(a["count"] for a in rep["actions"]), bold=True, fill=head, align="center")
    gr = 10
    for a in rep["actions"]:
        put(f"G{gr}", a["label"], wrap=True)
        put(f"H{gr}", a["count"], align="center")
        gr += 1
    put(f"G{gr}", "Total action taken to correct near miss errors", bold=True, fill=soft, wrap=True)
    put(f"H{gr}", sum(a["count"] for a in rep["actions"]), bold=True, fill=soft, align="center")
    gr += 2
    put(f"G{gr}", "Recommendations", bold=True, fill=head)
    put(f"H{gr}", sum(x["count"] for x in rep["recs"]), bold=True, fill=head, align="center")
    gr += 1
    for x in rep["recs"]:
        put(f"G{gr}", x["label"], wrap=True)
        put(f"H{gr}", x["count"], align="center")
        gr += 1
    actions_end = gr + 1

    # ---- drugs involved
    r = max(left_end, causes_end, actions_end) + 1
    put(f"E{r}", "Drug involved", bold=True, border=False)
    r += 1
    put(f"D{r}", "", fill=head)
    put(f"E{r}", "Drugs most frequently involved in near miss error", bold=True, fill=head)
    put(f"F{r}", "Frequency", bold=True, fill=head, align="center")
    note_row = r
    r += 1
    for i in range(20):
        d = rep["drugs"][i] if i < len(rep["drugs"]) else None
        put(f"D{r}", i + 1, align="center")
        put(f"E{r}", d["name"] if d else "", wrap=True)
        put(f"F{r}", d["count"] if d else "", align="center")
        r += 1
    notes = [
        "The names of the drugs that show repeated human errors must be listed under this section.",
        "Causes such as look-a-like, sound-a-like, wrong formulation packed, wrong strength packed.",
        "Time of day must be distinctly reported. No combining of lunch time with afternoon.",
        "Only the Summary Report must be submitted in the Monthly Report.",
    ]
    for i, n in enumerate(notes):
        ws.merge_cells(start_row=note_row + i * 2, start_column=7, end_row=note_row + i * 2 + 1, end_column=8)
        put(f"G{note_row + i * 2}", n, border=False, wrap=True, size=9)
        ws[f"G{note_row + i * 2}"].alignment = Alignment(wrap_text=True, vertical="top")

    for col, w in {"A": 50, "B": 9, "C": 3, "D": 5, "E": 58, "F": 11, "G": 50, "H": 10}.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---- log sheet, same columns as the Near Miss Form
    lg = wb.create_sheet("Near Miss Log")
    cols = ["No.", "Pharmacy", "Date", "Time of day", "Staffing level", "Personnel involved", "Reported by",
            "Name and brand of drug", "Type of pharmacy error", "Possible causes", "Things to consider",
            "Action taken", "Potential adverse event discussion", "Notes"]
    widths = [7, 22, 12, 14, 30, 30, 24, 46, 26, 46, 46, 40, 16, 40]
    lg.append(cols)
    for i, w in enumerate(widths):
        lg.column_dimensions[chr(65 + i)].width = w
    lbl = lambda pairs: dict(pairs)
    pers, rb = lbl(C.PERSONNEL), lbl(C.REPORTED_BY)
    cause, act = lbl(C.CAUSES), lbl(C.ACTIONS)
    tm = {c: n for c, n, _ in C.TIMES}
    tp = {c: f"{c} - {n}" for c, n, _ in C.ERROR_TYPES}
    for x in rep["_rows"]:
        staff = ", ".join(f"{x[k]} {l.lower()}" for k, l in C.STAFFING if x[k])
        names = f" ({x['personnel_names']})" if x["personnel_names"] else ""
        by = rb[x["reported_by_role"]] + (f" ({x['reported_by_name']})" if x["reported_by_name"] else "")
        actions = [act[a] if a != "other_action" or not x["action_other"] else f"Others: {x['action_other']}"
                   for a in x["actions"]]
        lg.append([
            x["id"], x["pharmacy_name"], dt.date.fromisoformat(x["date"]), tm[x["time_of_day"]], staff or "-",
            ", ".join(pers[p] for p in x["personnel"]) + names, by, x["drug_name"], tp[x["error_type"]],
            "; ".join(cause[c] for c in x["causes"]), "; ".join(x["considerations"]), "; ".join(actions),
            "Yes" if x["adverse_discussed"] else "", x["notes"],
        ])
    for c in lg[1]:
        c.font = Font(name="Arial", size=10, bold=True)
        c.fill = head
        c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in lg.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.border = box
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[2].number_format = "dd mmm yyyy"
    lg.freeze_panes = "A2"
    lg.auto_filter.ref = lg.dimensions
    lg.page_setup.orientation = "landscape"
    lg.page_setup.fitToWidth = 1
    lg.page_setup.fitToHeight = 0
    lg.sheet_properties.pageSetUpPr.fitToPage = True
    return wb


@app.get("/api/report.xlsx")
def report_xlsx():
    month, pid, e = report_args()
    if e:
        return err(e)
    rep = compute_report(month, pid)
    if rep is None:
        return err("Unknown pharmacy.", 404)
    buf = io.BytesIO()
    build_workbook(rep).save(buf)
    buf.seek(0)
    slug = re.sub(r"[^a-z0-9]+", "-", (rep["pharmacy"]["name"] if rep["pharmacy"] else "all-pharmacies").lower()).strip("-")
    return send_file(
        buf, as_attachment=True, download_name=f"near-miss-report_{month}_{slug}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=4400)
